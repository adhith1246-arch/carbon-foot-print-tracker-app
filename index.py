"""
Vercel Serverless Function — Campus Carbon Tracker API
Uses /tmp for JSON storage (Vercel's writable directory)
"""
import json, os, io, hashlib, secrets
from datetime import datetime, timedelta
from flask import Flask, request, jsonify, send_file, session
from flask_cors import CORS

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'greencampus-hackathon-2026-secret')
CORS(app, supports_credentials=True)

# Vercel uses /tmp for writable storage (ephemeral — resets on cold start)
DATA_DIR = '/tmp/carbon-data'
os.makedirs(DATA_DIR, exist_ok=True)
EMISSIONS_FILE = os.path.join(DATA_DIR, 'emissions.json')
USERS_FILE = os.path.join(DATA_DIR, 'users.json')

# ─── Helpers ───
FACTORS = {
    'electricity_kwh': 0.82, 'diesel_litre': 2.68, 'vehicle_km': 0.21,
    'veg_meal': 0.50, 'nonveg_meal': 3.00, 'food_waste_kg': 2.50, 'water_kl': 0.34,
}

def load_json(path):
    if os.path.exists(path):
        with open(path, 'r') as f:
            return json.load(f)
    return []

def save_json(path, data):
    with open(path, 'w') as f:
        json.dump(data, f, indent=2, default=str)

def hash_pw(pw):
    return hashlib.sha256(pw.encode()).hexdigest()

def calc_co2(d):
    e = d.get('electricity_kwh', 0) * FACTORS['electricity_kwh']
    di = d.get('diesel_litres', 0) * FACTORS['diesel_litre']
    v = d.get('vehicles_entered', 0) * 2 * FACTORS['vehicle_km']
    f = d.get('veg_meals', 0) * FACTORS['veg_meal'] + d.get('nonveg_meals', 0) * FACTORS['nonveg_meal']
    w = d.get('food_waste_kg', 0) * FACTORS['food_waste_kg']
    wa = d.get('water_kl', 0) * FACTORS['water_kl']
    total = e + di + v + f + w + wa
    return {k: round(val, 2) for k, val in
            {'electricity': e, 'diesel': di, 'vehicles': v, 'food': f, 'food_waste': w, 'water': wa, 'total': total}.items()}

def generate_suggestions(co2):
    tips = []
    if co2.get('electricity', 0) > 350:
        tips.append({'icon': 'lightbulb', 'type': 'warn', 'title': 'High Electricity',
                     'message': 'Switch off ACs in unoccupied rooms.',
                     'saving': f"~{((co2['electricity']-350)*0.3):.1f} kg CO2/day"})
    if co2.get('food_waste', 0) > 25:
        tips.append({'icon': 'food', 'type': 'warn', 'title': 'Excess Food Waste',
                     'message': 'Reduce preparation by 15-20%.',
                     'saving': f"~{(co2['food_waste']*0.4):.1f} kg CO2/day"})
    tips.append({'icon': 'sun', 'type': 'info', 'title': 'Switch to Solar',
                 'message': 'Solar panels could offset 60% of electricity emissions.',
                 'saving': f"~{(co2.get('electricity',0)*0.6):.1f} kg CO2/day"})
    return tips

# ─── Auth APIs ───
@app.route('/api/auth/register', methods=['POST'])
def register():
    data = request.json
    name = data.get('name', '').strip()
    email = data.get('email', '').strip().lower()
    password = data.get('password', '')
    department = data.get('department', '')
    if not all([name, email, password]):
        return jsonify({'success': False, 'error': 'All fields required'}), 400
    users = load_json(USERS_FILE)
    if any(u['email'] == email for u in users):
        return jsonify({'success': False, 'error': 'Email already registered'}), 400
    user = {'id': len(users)+1, 'name': name, 'email': email,
            'password': hash_pw(password), 'department': department,
            'role': 'admin' if len(users) == 0 else 'user',
            'created_at': datetime.now().isoformat()}
    users.append(user)
    save_json(USERS_FILE, users)
    session['user'] = {'id': user['id'], 'name': name, 'email': email,
                       'department': department, 'role': user['role']}
    return jsonify({'success': True, 'user': session['user']})

@app.route('/api/auth/login', methods=['POST'])
def login():
    data = request.json
    email = data.get('email', '').strip().lower()
    password = data.get('password', '')
    users = load_json(USERS_FILE)
    user = next((u for u in users if u['email'] == email and u['password'] == hash_pw(password)), None)
    if not user:
        return jsonify({'success': False, 'error': 'Invalid email or password'}), 401
    session['user'] = {'id': user['id'], 'name': user['name'], 'email': user['email'],
                       'department': user['department'], 'role': user['role']}
    return jsonify({'success': True, 'user': session['user']})

@app.route('/api/auth/logout', methods=['POST'])
def logout():
    session.pop('user', None)
    return jsonify({'success': True})

@app.route('/api/auth/me')
def me():
    if 'user' in session:
        return jsonify({'success': True, 'user': session['user']})
    return jsonify({'success': False}), 401

# ─── Emissions APIs ───
@app.route('/api/health')
def health():
    return jsonify({'status': 'OK', 'timestamp': datetime.now().isoformat()})

@app.route('/api/emissions/add', methods=['POST'])
def add_emission():
    data = request.json
    co2 = calc_co2(data)
    record = {
        'id': len(load_json(EMISSIONS_FILE)) + 1,
        'department': data.get('department', ''),
        'date': data.get('date', datetime.now().strftime('%Y-%m-%d')),
        'electricity_kwh': data.get('electricity_kwh', 0),
        'diesel_litres': data.get('diesel_litres', 0),
        'vehicles_entered': data.get('vehicles_entered', 0),
        'veg_meals': data.get('veg_meals', 0),
        'nonveg_meals': data.get('nonveg_meals', 0),
        'food_waste_kg': data.get('food_waste_kg', 0),
        'water_kl': data.get('water_kl', 0),
        'co2': co2,
        'submitted_by': session.get('user', {}).get('name', 'Anonymous'),
        'submitted_at': datetime.now().isoformat(),
    }
    records = load_json(EMISSIONS_FILE)
    records.append(record)
    save_json(EMISSIONS_FILE, records)
    return jsonify({'success': True, 'data': record})

@app.route('/api/emissions/today')
def today_emissions():
    today = datetime.now().strftime('%Y-%m-%d')
    records = [r for r in load_json(EMISSIONS_FILE) if r['date'] == today]
    totals = {'total':0,'electricity':0,'diesel':0,'vehicles':0,'food':0,'food_waste':0,'water':0}
    for r in records:
        for k in totals: totals[k] += r['co2'].get(k, 0)
    return jsonify({'success': True, 'date': today, 'totals': {k: round(v,2) for k,v in totals.items()}, 'count': len(records)})

@app.route('/api/emissions/trend')
def trend():
    days = int(request.args.get('days', 30))
    cutoff = (datetime.now() - timedelta(days=days)).strftime('%Y-%m-%d')
    by_date = {}
    for r in load_json(EMISSIONS_FILE):
        if r['date'] >= cutoff:
            by_date[r['date']] = by_date.get(r['date'], 0) + r['co2']['total']
    return jsonify({'success': True, 'trend': [{'date': k, 'co2': round(v,2)} for k,v in sorted(by_date.items())]})

@app.route('/api/emissions/recent')
def recent():
    records = sorted(load_json(EMISSIONS_FILE), key=lambda r: r.get('submitted_at',''), reverse=True)[:10]
    return jsonify({'success': True, 'records': records})

@app.route('/api/leaderboard')
def leaderboard():
    cutoff = (datetime.now() - timedelta(days=7)).strftime('%Y-%m-%d')
    dept_map = {}
    for r in load_json(EMISSIONS_FILE):
        if r['date'] >= cutoff:
            d = r['department']
            if d not in dept_map: dept_map[d] = {'total':0,'entries':0}
            dept_map[d]['total'] += r['co2']['total']
            dept_map[d]['entries'] += 1
    lb = sorted([{'department':n,'total_co2':round(d['total'],1),
                  'avg_daily':round(d['total']/max(d['entries'],1),1),'entries':d['entries']}
                 for n,d in dept_map.items()], key=lambda x: x['total_co2'])
    return jsonify({'success': True, 'leaderboard': lb})

@app.route('/api/suggestions')
def suggestions():
    today = datetime.now().strftime('%Y-%m-%d')
    co2 = {'electricity':0,'diesel':0,'vehicles':0,'food':0,'food_waste':0}
    for r in load_json(EMISSIONS_FILE):
        if r['date'] == today:
            for k in co2: co2[k] += r['co2'].get(k, 0)
    return jsonify({'success': True, 'suggestions': generate_suggestions(co2)})

# ─── Excel Export ───
@app.route('/api/export/excel')
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    records = load_json(EMISSIONS_FILE)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Emissions Data'
    hf = Font(bold=True, color='FFFFFF', size=11)
    hfill = PatternFill(start_color='0F5132', end_color='0F5132', fill_type='solid')
    tb = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
    headers = ['ID','Date','Department','Electricity (kWh)','Diesel (L)','Vehicles',
               'Veg Meals','Non-Veg Meals','Food Waste (kg)','Water (kL)',
               'CO2 Total (kg)','Submitted By','Submitted At']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hf; c.fill = hfill; c.alignment = Alignment(horizontal='center'); c.border = tb
    for i, r in enumerate(records, 2):
        row = [r.get('id'), r.get('date'), r.get('department'),
               r.get('electricity_kwh',0), r.get('diesel_litres',0), r.get('vehicles_entered',0),
               r.get('veg_meals',0), r.get('nonveg_meals',0), r.get('food_waste_kg',0), r.get('water_kl',0),
               r.get('co2',{}).get('total',0), r.get('submitted_by',''), r.get('submitted_at','')]
        for col, val in enumerate(row, 1):
            c = ws.cell(row=i, column=col, value=val); c.border = tb
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or '')) for c in col)+3, 25)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name=f"carbon_report_{datetime.now().strftime('%Y%m%d')}.xlsx",
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ─── Sensor API ───
@app.route('/api/sensors/push', methods=['POST'])
def sensor_push():
    data = request.json
    co2 = calc_co2(data)
    record = {
        'id': len(load_json(EMISSIONS_FILE)) + 1,
        'department': data.get('department', 'Unknown'),
        'date': datetime.now().strftime('%Y-%m-%d'),
        'electricity_kwh': data.get('electricity_kwh', 0),
        'diesel_litres': data.get('diesel_litres', 0),
        'vehicles_entered': data.get('vehicles_entered', 0),
        'veg_meals': data.get('veg_meals', 0),
        'nonveg_meals': data.get('nonveg_meals', 0),
        'food_waste_kg': data.get('food_waste_kg', 0),
        'water_kl': data.get('water_kl', 0),
        'co2': co2, 'source': 'sensor',
        'sensor_id': data.get('sensor_id', 'unknown'),
        'submitted_by': 'IoT Sensor',
        'submitted_at': datetime.now().isoformat(),
    }
    records = load_json(EMISSIONS_FILE)
    records.append(record)
    save_json(EMISSIONS_FILE, records)
    return jsonify({'success': True, 'data': record})

# Vercel handler
app = app
